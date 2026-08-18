"""Extract text/metadata from documents, images, audio, video."""

from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
from pathlib import Path

from .config import load_dotenv, llm_config
from .media_types import AUDIO_EXT, DOCUMENT_EXT, IMAGE_EXT, VIDEO_EXT

# PDF OCR 默认：页文本过少则整页扫描 OCR；内嵌图最小边长过滤图标
_PDF_OCR_MIN_PAGE_CHARS = 40
_PDF_OCR_MIN_IMG_SIDE = 80
_PDF_OCR_MAX_PAGES = 80
_PDF_OCR_MAX_IMAGES_PER_PAGE = 6
_PDF_OCR_DEFAULT_DPI = 160


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return _read_text(path)
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".doc":
        return _extract_doc(path)
    if ext == ".pptx":
        return _extract_pptx(path)
    if ext == ".ppt":
        return _extract_ppt(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".xls":
        return _extract_xls(path)
    if ext in IMAGE_EXT:
        return _extract_image(path)
    if ext in AUDIO_EXT:
        return _extract_audio(path)
    if ext in VIDEO_EXT:
        return _extract_video(path)
    return f"[不支持的格式] {path.name}"


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb2312"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, OSError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _tesseract_available() -> bool:
    return bool(shutil.which("tesseract"))


def _pymupdf_available() -> bool:
    try:
        import fitz  # noqa: F401  # type: ignore

        return True
    except ImportError:
        return False


def _pdf_ocr_enabled() -> bool:
    """MYKNOWLEDGE_PDF_OCR 默认开启；显式 0/false/off 关闭。"""
    v = (os.environ.get("MYKNOWLEDGE_PDF_OCR") or "1").strip().lower()
    return v not in ("0", "false", "no", "off")


def _pdf_ocr_dpi() -> int:
    raw = (os.environ.get("MYKNOWLEDGE_PDF_OCR_DPI") or "").strip()
    try:
        dpi = int(raw) if raw else _PDF_OCR_DEFAULT_DPI
    except ValueError:
        dpi = _PDF_OCR_DEFAULT_DPI
    return max(72, min(300, dpi))


def _ocr_pil(img) -> str:
    """OCR a PIL image; empty string on failure / missing tesseract."""
    if not _tesseract_available():
        return ""
    try:
        import pytesseract  # type: ignore

        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
        return (text or "").strip()
    except Exception:
        return ""


def _pdf_caption_enabled() -> bool:
    """MYKNOWLEDGE_PDF_CAPTION 默认关；1/true/on 启用图表视觉图注。"""
    v = (os.environ.get("MYKNOWLEDGE_PDF_CAPTION") or "0").strip().lower()
    return v not in ("0", "false", "no", "off", "")


def _blocks_to_multicolumn_text(
    blocks: list[tuple[float, float, float, float, str]], page_width: float
) -> str:
    """按分栏阅读序：先左栏自上而下，再右栏。单栏则按 y/x 排序。"""
    if not blocks:
        return ""
    gap_thr = max(24.0, page_width * 0.12)
    mids = sorted((b[0] + b[2]) / 2 for b in blocks)
    best_gap = 0.0
    split_x: float | None = None
    for i in range(len(mids) - 1):
        g = mids[i + 1] - mids[i]
        if g > best_gap and g >= gap_thr:
            best_gap = g
            split_x = (mids[i] + mids[i + 1]) / 2

    def _col_join(col: list[tuple[float, float, float, float, str]]) -> str:
        ordered = sorted(col, key=lambda b: (round(b[1] / 6) * 6, b[0]))
        return "\n".join(b[4] for b in ordered if (b[4] or "").strip())

    if split_x is None:
        return _col_join(blocks)
    left = [b for b in blocks if (b[0] + b[2]) / 2 < split_x]
    right = [b for b in blocks if (b[0] + b[2]) / 2 >= split_x]
    # 需两侧都有实质内容才当双栏，避免误判
    if len(left) < 2 or len(right) < 2:
        return _col_join(blocks)
    parts = [_col_join(left), _col_join(right)]
    return "\n\n".join(p for p in parts if p)


def _pdf_page_texts_pymupdf(path: Path) -> list[str] | None:
    """PyMuPDF 分栏阅读序提取；失败返回 None。"""
    if not _pymupdf_available():
        return None
    try:
        import fitz  # type: ignore

        pages: list[str] = []
        with fitz.open(str(path)) as doc:
            for page in doc:
                d = page.get_text("dict") or {}
                blocks: list[tuple[float, float, float, float, str]] = []
                for b in d.get("blocks") or []:
                    if b.get("type") != 0:
                        continue
                    line_parts: list[str] = []
                    for line in b.get("lines") or []:
                        spans = "".join(
                            (s.get("text") or "") for s in (line.get("spans") or [])
                        )
                        if spans.strip():
                            line_parts.append(spans)
                    text = "\n".join(line_parts).strip()
                    if not text:
                        continue
                    bbox = b.get("bbox") or (0, 0, 0, 0)
                    x0, y0, x1, y1 = (float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3]))
                    blocks.append((x0, y0, x1, y1, text))
                pages.append(_blocks_to_multicolumn_text(blocks, float(page.rect.width)))
        return pages
    except Exception:
        return None


def _vision_caption_image(img, ocr_hint: str = "") -> str:
    """可选视觉图注（需 MYKNOWLEDGE_PDF_CAPTION=1 且 LLM 支持 image_url）。失败返回空。"""
    if not _pdf_caption_enabled():
        return ""
    cfg = llm_config()
    if not cfg.get("api_key") or not cfg.get("chat_url"):
        return ""
    try:
        import base64
        import json
        import urllib.error
        import urllib.request

        from PIL import Image  # type: ignore

        # 缩小以控制请求体积
        work = img.convert("RGB")
        w, h = work.size
        max_side = 768
        if max(w, h) > max_side:
            scale = max_side / float(max(w, h))
            work = work.resize((max(1, int(w * scale)), max(1, int(h * scale))))
        buf = io.BytesIO()
        work.save(buf, format="JPEG", quality=75)
        b64 = base64.b64encode(buf.getvalue()).decode("ascii")
        hint = (ocr_hint or "").strip()[:400]
        prompt = "用一句中文概括此图作为图注，勿编造细节；无法识别则只返回空字符串。"
        if hint:
            prompt += f"\n图内 OCR 参考（可忽略错误）：{hint}"
        payload = {
            "model": cfg["model"],
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:image/jpeg;base64,{b64}"},
                        },
                    ],
                }
            ],
            "temperature": 0.2,
            "max_tokens": 80,
            "stream": False,
        }
        req = urllib.request.Request(
            cfg["chat_url"],
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {cfg['api_key']}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=45) as resp:
            obj = json.loads(resp.read().decode("utf-8", errors="replace"))
        text = (
            ((obj.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
        ).strip()
        if not text or text in ("空", "无", "（空）", "N/A", "n/a"):
            return ""
        return text[:200]
    except Exception:
        return ""


def _pdf_page_texts(path: Path) -> list[str] | None:
    """Return per-page text layer strings (may be empty). None = open failed."""
    ordered = _pdf_page_texts_pymupdf(path)
    if ordered is not None:
        return ordered

    try:
        import pdfplumber  # type: ignore

        pages: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                pages.append((page.extract_text() or "").strip())
        return pages
    except ImportError:
        pass
    except Exception:
        pass

    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        return [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception:
        return None


def _ocr_embedded_images(doc, page, page_no: int) -> list[str]:
    """OCR large enough embedded images on one PyMuPDF page."""
    from PIL import Image  # type: ignore

    blocks: list[str] = []
    try:
        images = page.get_images(full=True) or []
    except Exception:
        return blocks
    seen: set[int] = set()
    img_i = 0
    caption_budget = 2 if _pdf_caption_enabled() else 0
    for info in images:
        if img_i >= _PDF_OCR_MAX_IMAGES_PER_PAGE:
            break
        xref = int(info[0])
        if xref in seen:
            continue
        seen.add(xref)
        try:
            raw = doc.extract_image(xref)
        except Exception:
            continue
        if not raw or not raw.get("image"):
            continue
        try:
            img = Image.open(io.BytesIO(raw["image"]))
            img.load()
        except Exception:
            continue
        w, h = img.size
        if w < _PDF_OCR_MIN_IMG_SIDE or h < _PDF_OCR_MIN_IMG_SIDE:
            continue
        ocr = _ocr_pil(img)
        if not ocr and caption_budget <= 0:
            continue
        img_i += 1
        if ocr:
            blocks.append(f"--- 第{page_no}页-图{img_i} OCR ---\n{ocr}")
        if caption_budget > 0:
            cap = _vision_caption_image(img, ocr)
            caption_budget -= 1
            if cap:
                blocks.append(f"--- 第{page_no}页-图{img_i} 图注 ---\n{cap}")
    return blocks


def _ocr_page_scan(page, page_no: int, dpi: int) -> str:
    """Rasterize entire page and OCR (scanned / weak text pages)."""
    from PIL import Image  # type: ignore

    try:
        import fitz  # type: ignore

        zoom = dpi / 72.0
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    except Exception:
        return ""
    ocr = _ocr_pil(img)
    if not ocr:
        return ""
    return f"--- 第{page_no}页-扫描 OCR ---\n{ocr}"


def _extract_pdf_ocr_augment(path: Path, page_texts: list[str]) -> tuple[list[str], list[str]]:
    """
    Return (per_page_extra_blocks_joined_as_list_of_strings, notices).
    Each entry in extras aligns with page_texts index (may be empty string).
    """
    notices: list[str] = []
    extras = [""] * len(page_texts)
    if not _pdf_ocr_enabled():
        return extras, notices
    if not _tesseract_available():
        notices.append("（未安装 Tesseract OCR，已跳过 PDF 内嵌图/扫描页识别；安装后重新导入可补全）")
        return extras, notices
    if not _pymupdf_available():
        notices.append("（未安装 PyMuPDF，已跳过 PDF 图/扫描 OCR；pip 安装 pymupdf 后重新导入可补全）")
        return extras, notices

    try:
        import fitz  # type: ignore

        doc = fitz.open(str(path))
    except Exception as e:
        notices.append(f"（PDF OCR 打开失败: {e}）")
        return extras, notices

    dpi = _pdf_ocr_dpi()
    n_pages = min(len(page_texts), doc.page_count, _PDF_OCR_MAX_PAGES)
    try:
        for i in range(n_pages):
            page = doc.load_page(i)
            page_no = i + 1
            text = page_texts[i] if i < len(page_texts) else ""
            chunks: list[str] = []
            if len(text.strip()) < _PDF_OCR_MIN_PAGE_CHARS:
                scan = _ocr_page_scan(page, page_no, dpi)
                if scan:
                    chunks.append(scan)
                else:
                    chunks.extend(_ocr_embedded_images(doc, page, page_no))
            else:
                chunks.extend(_ocr_embedded_images(doc, page, page_no))
            if chunks:
                extras[i] = "\n\n".join(chunks)
    finally:
        try:
            doc.close()
        except Exception:
            pass
    return extras, notices


def _extract_pdf(path: Path) -> str:
    page_texts = _pdf_page_texts(path)
    if page_texts is None:
        # 两路文本库都失败时仍尝试纯 OCR
        page_texts = []
        if _pymupdf_available():
            try:
                import fitz  # type: ignore

                with fitz.open(str(path)) as doc:
                    page_texts = [""] * doc.page_count
            except Exception as e:
                return f"[PDF 提取失败] {path.name}: {e}"
        else:
            return f"[PDF 提取失败] {path.name}"

    if not page_texts:
        return f"[PDF 无页面] {path.name}"

    from .pdf_tables import extract_tables_markdown, format_table_blocks

    table_by_page = format_table_blocks(extract_tables_markdown(path))
    pages_with_digital_tables = set(table_by_page.keys())

    layout_by_page: dict[int, list[str]] = {}
    layout_notices: list[str] = []
    try:
        from .pdf_layout import extract_layout_table_blocks

        layout_by_page, layout_notices = extract_layout_table_blocks(
            path, page_texts, skip_pages=pages_with_digital_tables
        )
    except Exception as e:
        layout_notices.append(f"（PDF 版面表格增强跳过: {e}）")

    extras, notices = _extract_pdf_ocr_augment(path, page_texts)
    notices.extend(layout_notices)
    parts: list[str] = []
    for i, text in enumerate(page_texts):
        page_no = i + 1
        body: list[str] = []
        if text.strip():
            body.append(f"--- 第{page_no}页 ---\n{text.strip()}")
        for tb in table_by_page.get(page_no, []):
            if not body:
                body.append(f"--- 第{page_no}页 ---\n")
            body.append(tb)
        for lb in layout_by_page.get(page_no, []):
            if not body:
                body.append(f"--- 第{page_no}页 ---\n")
            body.append(lb)
        extra = extras[i] if i < len(extras) else ""
        if extra:
            if not body:
                # 弱文本页仅有 OCR：仍放页标方便 RAG
                body.append(f"--- 第{page_no}页 ---\n")
            body.append(extra)
        if body:
            parts.append("\n\n".join(body).strip())

    if parts:
        out = "\n\n".join(parts)
        if notices:
            out += "\n\n" + "\n".join(notices)
        return _maybe_augment_pdf_llm_parse(path, out, notices)

    msg = f"[PDF 无文本层] {path.name}"
    if notices:
        msg += "\n" + "\n".join(notices)
    elif not _pdf_ocr_enabled():
        msg += "\n（PDF OCR 已关闭：MYKNOWLEDGE_PDF_OCR=0）"
    # 无文本层时若开启 LLM 版面解析，直接用视觉结果
    try:
        from .pdf_llm_parse import pdf_llm_parse_enabled, parse_pdf_with_llm

        if pdf_llm_parse_enabled():
            r = parse_pdf_with_llm(path, merge_table=True)
            if r.get("ok") and (r.get("text") or "").strip():
                return (r["text"] or "") + "\n\n（本文件由 MYKNOWLEDGE_PDF_LLM_PARSE 视觉解析）"
    except Exception:
        pass
    return msg


def _maybe_augment_pdf_llm_parse(path: Path, out: str, notices: list[str]) -> str:
    """MYKNOWLEDGE_PDF_LLM_PARSE=1 时用视觉解析覆盖/增强（默认关，费 token）。"""
    try:
        from .pdf_llm_parse import pdf_llm_parse_enabled, parse_pdf_with_llm
    except Exception:
        return out
    if not pdf_llm_parse_enabled():
        return out
    # 仅当本地提取偏弱或显式 force 时替换，避免每份 PDF 都烧 token
    force = (os.environ.get("MYKNOWLEDGE_PDF_LLM_PARSE_FORCE") or "").strip().lower() in (
        "1",
        "true",
        "yes",
        "on",
    )
    weak = len(re.sub(r"\s+", "", out)) < 200 or any("无文本" in n for n in notices)
    if not force and not weak:
        return out + "\n\n（已跳过 LLM 版面解析：本地提取足够；设 MYKNOWLEDGE_PDF_LLM_PARSE_FORCE=1 可强制）"
    try:
        r = parse_pdf_with_llm(path, merge_table=True)
        if r.get("ok") and (r.get("text") or "").strip():
            return (
                (r["text"] or "").strip()
                + "\n\n（本文件由 MYKNOWLEDGE_PDF_LLM_PARSE 视觉解析；本地提取已旁路）"
            )
    except Exception as e:
        return out + f"\n\n（LLM 版面解析失败: {e}）"
    return out



def _extract_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts) if parts else f"[DOCX 无文本] {path.name}"
    except Exception as e:
        return f"[DOCX 提取失败] {path.name}: {e}"


def _extract_doc(path: Path) -> str:
    # Legacy .doc: try LibreOffice headless, else placeholder
    for cmd in (
        ["soffice", "--headless", "--convert-to", "txt:Text", "--outdir", str(path.parent), str(path)],
        ["libreoffice", "--headless", "--convert-to", "txt:Text", "--outdir", str(path.parent), str(path)],
    ):
        try:
            subprocess.run(cmd, capture_output=True, timeout=120, check=False)
            txt = path.with_suffix(".txt")
            if txt.is_file():
                content = _read_text(txt)
                try:
                    txt.unlink()
                except OSError:
                    pass
                if content.strip():
                    return content
        except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
            continue
    return (
        f"[.doc 旧格式] {path.name}\n"
        "未能自动提取文本。请转换为 docx 后重新导入，或在知识库条目中手动补充摘要。"
    )


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore

        prs = Presentation(str(path))
        parts: list[str] = []
        for si, slide in enumerate(prs.slides, 1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                parts.append(f"--- 幻灯片 {si} ---\n" + "\n".join(slide_text))
        return "\n\n".join(parts) if parts else f"[PPTX 无文本] {path.name}"
    except Exception as e:
        return f"[PPTX 提取失败] {path.name}: {e}"


def _extract_ppt(path: Path) -> str:
    return (
        f"[.ppt 旧格式] {path.name}\n"
        "请另存为 pptx 后重新导入，或使用 LibreOffice 转换。"
    )


def _extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        for name in wb.sheetnames:
            ws = wb[name]
            rows: list[str] = []
            for row in ws.iter_rows(max_row=500, values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"--- 工作表: {name} ---\n" + "\n".join(rows[:200]))
        wb.close()
        return "\n\n".join(parts) if parts else f"[XLSX 无数据] {path.name}"
    except Exception as e:
        return f"[XLSX 提取失败] {path.name}: {e}"


def _extract_xls(path: Path) -> str:
    try:
        import xlrd  # type: ignore

        book = xlrd.open_workbook(str(path))
        parts: list[str] = []
        for si in range(book.nsheets):
            sh = book.sheet_by_index(si)
            rows: list[str] = []
            for ri in range(min(sh.nrows, 500)):
                cells = [str(sh.cell_value(ri, ci)).strip() for ci in range(sh.ncols)]
                cells = [c for c in cells if c]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"--- 工作表: {sh.name} ---\n" + "\n".join(rows))
        return "\n\n".join(parts) if parts else f"[XLS 无数据] {path.name}"
    except Exception as e:
        return f"[XLS 提取失败] {path.name}: {e}"


def _extract_image(path: Path) -> str:
    parts = [f"图片文件: {path.name}"]
    ocr = ""
    try:
        from PIL import Image  # type: ignore

        with Image.open(path) as img:
            parts.append(f"尺寸: {img.size[0]}x{img.size[1]}")
            ocr = _ocr_pil(img.copy())
    except Exception:
        pass
    if ocr:
        parts.append("--- OCR 识别 ---\n" + ocr)
    elif _tesseract_available():
        parts.append("（OCR 未识别到文字，可在 GUI 中补充描述）")
    else:
        parts.append("（未安装 Tesseract OCR，仅索引文件名与尺寸；安装后可识别图中文字）")
    return "\n".join(parts)


def _media_tags(path: Path) -> str:
    try:
        from mutagen import File as MutagenFile  # type: ignore

        meta = MutagenFile(str(path))
        if meta is None:
            return ""
        lines: list[str] = []
        for key in ("title", "artist", "album", "date", "genre"):
            val = meta.get(key)
            if val:
                if isinstance(val, list):
                    val = val[0]
                lines.append(f"{key}: {val}")
        if hasattr(meta, "info") and meta.info:
            info = meta.info
            if hasattr(info, "length") and info.length:
                lines.append(f"时长: {int(info.length)}s")
        return "\n".join(lines)
    except Exception:
        return ""


def _transcribe_whisper(path: Path) -> str:
    load_dotenv()
    cfg = llm_config()
    if not cfg["api_key"]:
        return ""
    try:
        import json
        import urllib.request

        base = cfg["chat_url"].replace("/chat/completions", "")
        if not base.endswith("/v1"):
            base = base.rstrip("/") + "/v1"
        url = f"{base}/audio/transcriptions"
        boundary = "----MyknowledgeBoundary"
        data = path.read_bytes()
        body = (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="model"\r\n\r\nwhisper-1\r\n'
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
            f"Content-Type: application/octet-stream\r\n\r\n"
        ).encode("utf-8") + data + f"\r\n--{boundary}--\r\n".encode("utf-8")
        req = urllib.request.Request(
            url,
            data=body,
            headers={
                "Authorization": f"Bearer {cfg['api_key']}",
                "Content-Type": f"multipart/form-data; boundary={boundary}",
            },
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=300) as resp:
            result = json.loads(resp.read().decode("utf-8"))
        return (result.get("text") or "").strip()
    except Exception:
        return ""


def _extract_audio(path: Path) -> str:
    parts = [f"音频文件: {path.name}"]
    tags = _media_tags(path)
    if tags:
        parts.append("--- 元数据 ---\n" + tags)
    transcript = _transcribe_whisper(path)
    if transcript:
        parts.append("--- 语音转写 ---\n" + transcript)
    else:
        parts.append("（未配置 Whisper 兼容 API 时仅索引元数据；可在 .env 配置 OPENAI_API_KEY 启用转写）")
    return "\n".join(parts)


def _extract_video(path: Path) -> str:
    parts = [f"视频文件: {path.name}"]
    tags = _media_tags(path)
    if tags:
        parts.append("--- 元数据 ---\n" + tags)
    # Extract audio track to temp mp3 via ffmpeg if available, then whisper
    audio_tmp = path.parent / f"._{path.stem}_audio.mp3"
    from .runtime_tools import find_ffmpeg

    ffmpeg = find_ffmpeg()
    if ffmpeg:
        try:
            subprocess.run(
                [ffmpeg, "-y", "-i", str(path), "-vn", "-acodec", "libmp3lame", str(audio_tmp)],
                capture_output=True,
                timeout=300,
                check=False,
            )
            if audio_tmp.is_file() and audio_tmp.stat().st_size > 0:
                parts.append(_extract_audio(audio_tmp))
                try:
                    audio_tmp.unlink()
                except OSError:
                    pass
                return "\n".join(parts)
        except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
            pass
    parts.append("（安装 ffmpeg 并配置 Whisper API 可自动转写视频音轨）")
    return "\n".join(parts)
