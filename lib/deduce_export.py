"""Export deduce report to Word (.docx) with graph image, header and page numbers."""

from __future__ import annotations

import base64
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from .produce_export import _add_runs_with_bold


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def _apply_yizhi_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.text = "易知"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        fp = footer.add_paragraph()
        _add_page_number(fp)


def _append_markdown_body(doc: Document, markdown: str) -> None:
    for line in (markdown or "").splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:].strip(), level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=3)
        elif re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(p, re.sub(r"^\d+\.\s+", "", s, count=1))
        else:
            p = doc.add_paragraph()
            _add_runs_with_bold(p, s)


def deduce_to_docx_bytes(
    markdown: str,
    *,
    title: str = "推演",
    graph_png_b64: str = "",
) -> bytes:
    doc = Document()
    _apply_yizhi_header_footer(doc)
    doc.add_heading(title or "推演", level=0)

    png = (graph_png_b64 or "").strip()
    if png.startswith("data:"):
        png = png.split(",", 1)[-1]
    if png:
        doc.add_heading("关系图", level=2)
        try:
            img = BytesIO(base64.b64decode(png))
            doc.add_picture(img, width=Inches(5.8))
        except Exception:
            doc.add_paragraph("（关系图导出失败）")

    _append_markdown_body(doc, markdown)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
