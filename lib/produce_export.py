"""Export produce Markdown to Word (.docx)."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document


def _add_runs_with_bold(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def markdown_to_docx_bytes(markdown: str, title: str = "产出") -> bytes:
    doc = Document()
    doc.add_heading(title or "产出", level=0)
    for line in (markdown or "").splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:].strip(), level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=3)
        elif re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(p, re.sub(r"^\d+\.\s+", "", s, count=1))
        else:
            p = doc.add_paragraph()
            _add_runs_with_bold(p, s)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
